"""Real nested DOCX and mixed native/raster PDFs through local OCR/model CLI."""
from pathlib import Path
from io import BytesIO
import json
import subprocess
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from billing_agent.__main__ import dump


def mixed_pdf(path,mode):
    from PIL import Image,ImageDraw,ImageFont
    from pypdf import PdfReader,PdfWriter
    from pypdf.generic import NameObject,DictionaryObject,NumberObject,DecodedStreamObject
    from scripts.run_linking_regression import pdf
    lines=['INVOICE No: COMPLEX-1001','Invoice Date: 2026-10-22','ISSUER: Alpha LLC','BILL TO: Beta LLC']
    if mode in {'mixed','same_amount'}:lines.append('Handling Fee USD 40.00')
    pdf(path,[lines])
    writer=PdfWriter();writer.append(PdfReader(path));page=writer.pages[0]
    image=Image.new('RGB',(1000,300),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48)
    text='Doc Fee USD 20.00' if mode=='mixed' else 'Handling Fee USD 50.00' if mode=='conflict' else 'Handling Fee USD 40.00'
    draw.text((50,60),text,font=font,fill='black')
    stream=DecodedStreamObject();stream.set_data(image.tobytes())
    stream.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),
                   NameObject('/Width'):NumberObject(1000),NameObject('/Height'):NumberObject(300),
                   NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
    page['/Resources'][NameObject('/XObject')]=DictionaryObject({NameObject('/Scan'):writer._add_object(stream.flate_encode())})
    content=page.get_contents().get_data()+b'\nq 500 0 0 150 40 150 cm /Scan Do Q\n'
    if mode in {'overlay','conflict'}:
        content+=b'BT /F1 24 Tf 3 Tr 65 246 Td (Handling Fee USD 40.00) Tj ET\n'
    stream=DecodedStreamObject();stream.set_data(content);page[NameObject('/Contents')]=writer._add_object(stream)
    output=BytesIO();writer.write(output);path.write_bytes(output.getvalue())


def main():
    root=ROOT/'outputs/complex-regression';root.mkdir(parents=True,exist_ok=True)
    fixtures=root/'fixtures';fixtures.mkdir(exist_ok=True)
    from docx import Document
    doc=Document();table=doc.add_table(rows=1,cols=2);cell=table.cell(0,0).merge(table.cell(0,1));cell.text='Nested test invoice'
    child=cell.add_table(rows=1,cols=1);leaf=child.cell(0,0).add_table(rows=3,cols=3)
    for row,values in zip(leaf.rows,[['Description','Currency','Amount'],['Handling Fee','USD','40'],['Doc Fee','USD','20']]):
        for c,value in zip(row.cells,values):c.text=value
    doc.save(fixtures/'nested.docx')
    cases=[('nested.docx',[20,40])]
    for mode,amounts in [('mixed',[20,40]),('same_amount',[40,40]),('overlay',[40]),('conflict',[40])]:
        mixed_pdf(fixtures/(mode+'.pdf'),mode);cases.append((mode+'.pdf',amounts))
    checks=[]
    for filename,expected in cases:
        output=root/'runs'/filename.replace('.','-')
        run=subprocess.run([sys.executable,str(ROOT/'scripts/run_local.py'),'--input',str(fixtures/filename),
                '--entity','Beta LLC','--pdf-mode','local','--image-mode','local','--local-ocr-cache',str(root/'ocr-cache'),
                '--extractor','llm','--output',str(output)],cwd=ROOT,capture_output=True,timeout=600)
        rows=json.loads((output/'result.json').read_text(encoding='utf-8'))
        report=json.loads((output/'run_report.json').read_text(encoding='utf-8'))
        amounts=sorted(r['charge_details']['amount'] for r in rows)
        passed=run.returncode==0 and amounts==expected and report['output_validation']['passed']
        if filename=='conflict.pdf':passed &= all(r['operation']=='PENDING_REVIEW' for r in rows)
        checks.append({'fixture':filename,'expected':expected,'actual':amounts,'passed':bool(passed),'report':str(output/'run_report.json')})
        print(filename,'PASS' if passed else 'FAIL',amounts,flush=True)
        dump(root/'validation.json',{'scope':'Synthetic nested documents and mixed PDFs, actual local OCR/model; no uploads',
                'checks':checks,'passed':len(checks)==5 and all(c['passed'] for c in checks)})
    return 0 if all(c['passed'] for c in checks) else 1


if __name__=='__main__':raise SystemExit(main())
