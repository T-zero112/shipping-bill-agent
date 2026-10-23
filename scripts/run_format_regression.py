"""Generate synthetic formats, then run the real local-model CLI on each input."""
import argparse
from email.message import EmailMessage
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import zipfile

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from billing_agent.conversion import convert
from billing_agent.__main__ import dump


def fixtures(root):
    import openpyxl
    from docx import Document
    from PIL import Image,ImageDraw,ImageFont
    root.mkdir(parents=True,exist_ok=True)
    doc=Document()
    doc.add_paragraph('SYNTHETIC INVOICE TEST-FMT-1001')
    table=doc.add_table(rows=3,cols=3)
    table.cell(0,0).merge(table.cell(0,2)).text='Synthetic test data only'
    for cell,text in zip(table.rows[1].cells,['Description','Currency','Amount']):cell.text=text
    for cell,text in zip(table.rows[2].cells,['Handling Fee','USD','40']):cell.text=text
    stream=BytesIO();doc.save(stream)
    (root/'table.docx').write_bytes(stream.getvalue())
    (root/'legacy.doc').write_bytes(convert(stream.getvalue(),'.docx','doc:MS Word 97'))
    book=openpyxl.Workbook();sheet=book.active
    sheet.merge_cells('A1:E1');sheet['A1']='SYNTHETIC INVOICE TEST-XLSX-1001'
    sheet.append(['Description','Qty','Unit Price','Currency','Amount'])
    sheet.append(['Handling Fee',2,20,'USD',40])
    stream=BytesIO();book.save(stream);plain=stream.getvalue()
    (root/'merged.xlsx').write_bytes(plain)
    (root/'legacy.xls').write_bytes(convert(plain,'.xlsx','xls:MS Excel 97'))
    sheet['E3']='=B3*C3';stream=BytesIO();book.save(stream)
    formula=stream.getvalue();(root/'formula-missing.xlsx').write_bytes(formula)
    target=BytesIO()
    with zipfile.ZipFile(BytesIO(formula)) as src,zipfile.ZipFile(target,'w') as dst:
        for name in src.namelist():
            raw=src.read(name)
            if name=='xl/worksheets/sheet1.xml':raw=raw.replace(b'<f>B3*C3</f><v></v>',b'<f>B3*C3</f><v>40</v>')
            dst.writestr(name,raw)
    (root/'formula-cached.xlsx').write_bytes(target.getvalue())
    inner=EmailMessage();inner.set_content('Synthetic attachment for format testing.')
    inner.add_attachment((root/'table.docx').read_bytes(),maintype='application',subtype='octet-stream',filename='invoice.docx')
    outer=EmailMessage();outer.set_content('<p>Synthetic attached invoice.</p><script>USD 9999</script>',subtype='html')
    outer.add_attachment(inner)
    (root/'nested.eml').write_bytes(outer.as_bytes())
    image=Image.new('RGB',(1500,1050),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',38)
    for i,line in enumerate(['SYNTHETIC INVOICE TEST-SCAN-1001','Invoice Date: 2026-10-22',
                             'Handling Fee USD 40.00','Total Amount: USD 40.00','Synthetic test data only']):
        draw.text((60,70+i*130),line,font=font,fill='black')
    image.save(root/'scan-preview.png')
    image.save(root/'scan.pdf',format='PDF',resolution=144)
    image.save(root/'multipage.tiff',format='TIFF',save_all=True,append_images=[image])
    return [('DOCX','table.docx',1),('DOC','legacy.doc',1),('XLSX merged','merged.xlsx',1),
            ('XLS','legacy.xls',1),('XLSX missing formula cache','formula-missing.xlsx',0),
            ('XLSX cached formula','formula-cached.xlsx',1),('EML nested DOCX','nested.eml',1),
            ('Scanned PDF','scan.pdf',1),('TIFF multipage','multipage.tiff',2)]


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--generate-only',action='store_true');args=ap.parse_args()
    root=ROOT/'outputs'/'format-regression'
    cases=fixtures(root/'fixtures')
    if args.generate_only:return 0
    checks=[]
    for label,filename,count in cases:
        output=root/'runs'/filename.replace('.','-')
        command=[sys.executable,str(ROOT/'scripts/run_local.py'),'--input',str(root/'fixtures'/filename),
                 '--entity','Synthetic Beta LLC','--pdf-mode','local','--image-mode','local',
                 '--local-ocr-cache',str(root/'ocr-cache'),'--extractor','llm','--output',str(output)]
        run=subprocess.run(command,cwd=ROOT,capture_output=True,timeout=900)
        output.mkdir(parents=True,exist_ok=True)
        (output/'cli.log').write_bytes(run.stdout+run.stderr)
        report=json.loads((output/'run_report.json').read_text(encoding='utf-8')) if (output/'run_report.json').exists() else {}
        rows=json.loads((output/'result.json').read_text(encoding='utf-8')) if (output/'result.json').exists() else []
        values=[str(r['charge_details'].get('amount')) for r in rows]
        passed=(run.returncode==0 and len(rows)==count and all(v not in {'None','null'} and float(v)==40 for v in values)
                and not report.get('extraction_error') and not report.get('extraction_errors')
                and not any('error' in issue for issue in report.get('issues',[])))
        if 'formula' in label.lower():
            passed &= bool(report.get('issues')) and all(r['operation']=='PENDING_REVIEW' for r in rows)
        if label=='EML nested DOCX':
            passed &= all('::' in e['source'] for r in rows for e in r['audit_trail']['evidence'])
        if label=='Scanned PDF':
            passed &= any(e['method']=='local_pdf_ocr_unverified' for r in rows for e in r['audit_trail']['evidence'])
        check={'format':label,'fixture':filename,'expected_fees':count,'actual_fees':len(rows),
               'amounts':values,'passed':bool(passed),'run_report':str(output/'run_report.json'),
               'parser_errors':[i for i in report.get('issues',[]) if 'error' in i]}
        checks.append(check);print(label,'PASS' if passed else 'FAIL',values,flush=True)
        dump(root/'validation.json',{'dataset':'Generated synthetic documents; actual local OCR/model calls; no remote uploads',
                                     'checks':checks,'passed':len(checks)==len(cases) and all(c['passed'] for c in checks)})
    return 0 if all(c['passed'] for c in checks) else 1


if __name__=='__main__':raise SystemExit(main())
