from io import BytesIO
from billing_agent.parsers import Parser
from billing_agent.hybrid import reconcile,image_regions
from billing_agent.extraction import extract_local
from billing_agent.normalization import normalize_extraction


def box(text,x,y,width=100):
    return {'text':text,'score':.99,'x0':x,'x1':x+width,'top':y,'bottom':y+20,
            'polygon':[[x,y],[x+width,y],[x+width,y+20],[x,y+20]]}


def test_nested_merged_word_tables_emit_leaf_rows_once_and_keep_real_duplicates(monkeypatch):
    from docx import Document
    doc=Document();outer=doc.add_table(rows=1,cols=2)
    cell=outer.cell(0,0).merge(outer.cell(0,1))
    cell.text='Nested invoice'
    child=cell.add_table(rows=1,cols=1)
    leaf=child.cell(0,0).add_table(rows=3,cols=3)
    for c,t in zip(leaf.rows[0].cells,['Description','Currency','Amount']):c.text=t
    for row in leaf.rows[1:]:
        for c,t in zip(row.cells,['Handling Fee','USD','40']):c.text=t
    stream=BytesIO();doc.save(stream)
    parser=Parser();parser.parse(stream.getvalue(),'nested.docx')
    fees=[e for e in parser.evidence if (e.structured or {}).get('charge_cells')]
    assert len(fees)==2 and len({e.location for e in fees})==2
    assert all(e.structured['depth']==2 for e in fees)
    assert all('Handling Fee' not in e.text for e in parser.evidence if (e.structured or {}).get('depth',0)<2)
    ext=extract_local(fees)
    assert len(ext.charges)==2 and all(c.amount==40 for c in ext.charges)


def test_spatial_native_overlap_suppressed_but_same_text_elsewhere_kept():
    frame={'page':1,'boxes':[box('Handling Fee USD 40',0,0,300),box('Handling Fee USD 40',0,80,300)]}
    native=[{'text':'Handling Fee USD 40','x0':40,'x1':190,'top':100,'bottom':110}]
    filtered,suppressed,conflicts=reconcile(frame,native,[40,100,240,200],2)
    assert len(suppressed)==1 and len(filtered['boxes'])==1 and not conflicts
    assert filtered['boxes'][0]['top']==140


def test_overlapping_native_ocr_amount_conflict_not_silently_deduped():
    frame={'page':1,'boxes':[box('USD 50',0,0)]}
    native=[{'text':'USD 40','x0':40,'x1':90,'top':100,'bottom':110}]
    filtered,suppressed,conflicts=reconcile(frame,native,[40,100,240,200],2)
    assert len(conflicts)==1 and not suppressed and not filtered['boxes']


def test_decimal_point_or_sign_difference_never_counts_as_duplicate():
    from billing_agent.hybrid import same_text
    assert not same_text('USD 400','USD 40.0')
    assert not same_text('USD -40','USD 40')


def test_overlapping_images_merge_and_small_logos_skip():
    class Page:
        width=600;height=800
        images=[{'x0':30,'top':100,'x1':400,'bottom':500},
                {'x0':50,'top':200,'x1':300,'bottom':400},
                {'x0':10,'top':10,'x1':50,'bottom':30}]
    assert image_regions(Page())==[[30,100,400,500]]


def test_conflict_region_is_not_sent_to_model_and_parent_warns_rows(monkeypatch):
    from billing_agent.models import Evidence,Charge,Extraction
    page=Evidence(id='p',source='mixed.pdf',location='page:1',text='INVOICE No: INV1',
                  structured={'review_reasons':['Hybrid OCR conflict; manual check']})
    row=Evidence(id='r',source='mixed.pdf',location='page:1/charge-row:1',parent_id='p',text='Handling Fee USD 40')
    region=Evidence(id='x',source='mixed.pdf',location='page:1/region:1',text='Handling Fee USD 50',
                    structured={'kind':'hybrid_conflict','review_reasons':['Native and OCR disagree']})
    assert not extract_local([region]).charges
    ext=normalize_extraction(Extraction(charges=[Charge(charge_item='Handling Fee',amount=40,currency='USD',
                                                     amount_quote='USD 40',evidence_ids=['r'])]),[page,row,region])
    assert ext.charges[0].warnings


def test_duplicate_amount_headers_are_reviewed_not_arbitrarily_collapsed():
    from billing_agent.tabular import charge_cells,ambiguous_headers
    headers=['Description','Currency','Amount','Amount']
    assert ambiguous_headers(headers)
    assert charge_cells(headers,['Handling Fee','USD',40,50]) is None


def test_vertical_merge_does_not_visit_nested_child_twice():
    from docx import Document
    doc=Document();outer=doc.add_table(rows=2,cols=1)
    cell=outer.cell(0,0).merge(outer.cell(1,0))
    child=cell.add_table(rows=2,cols=3)
    for row,values in zip(child.rows,[['Description','Currency','Amount'],['Handling Fee','USD','40']]):
        for c,v in zip(row.cells,values):c.text=v
    stream=BytesIO();doc.save(stream)
    parser=Parser();parser.parse(stream.getvalue(),'vertical.docx')
    assert len([e for e in parser.evidence if (e.structured or {}).get('charge_cells')])==1


def test_identical_content_regions_at_distinct_positions_stay_distinct():
    class Page:
        width=600;height=800
        images=[{'x0':30,'top':100,'x1':400,'bottom':250},
                {'x0':30,'top':400,'x1':400,'bottom':550}]
    assert len(image_regions(Page()))==2
