from email.message import EmailMessage
from io import BytesIO
from billing_agent.parsers import Parser


def test_eml_body_and_attachment_recursion():
    message = EmailMessage()
    message['Subject'] = 'Synthetic charge'
    message.set_content('H/C USD 50')
    message.add_attachment(b'Doc Fee USD 10', maintype='application', subtype='octet-stream', filename='charge.txt')
    parser = Parser()
    parser.parse(message.as_bytes(), 'sample.eml')
    assert any(e.location == 'body' and 'USD 50' in e.text for e in parser.evidence)
    assert any('::charge.txt' in e.source for e in parser.evidence)
    statements = [e for e in parser.evidence if e.structured and e.structured.get('kind')=='email_statement']
    assert len(statements)==1
    assert statements[0].text=='H/C USD 50'
    assert statements[0].parent_id == next(e.id for e in parser.evidence if e.location=='body')


def test_duplicate_file_and_unsupported_are_visible():
    parser = Parser()
    parser.parse(b'Hello', 'a.txt'); parser.parse(b'Hello', 'b.txt')
    parser.parse(b'legacy', 'test.doc')
    assert len(parser.evidence) == 1
    assert len(parser.issues) == 2


def test_image_without_url_is_not_silently_skipped():
    parser = Parser()
    parser.parse(b'fake image', 'scan.png')
    assert parser.issues[0]['error'].startswith('ValueError: OCR needs')


def test_xlsx_docx_tables():
    import openpyxl
    from docx import Document
    book = openpyxl.Workbook(); book.active.append(['Fee','Amount']); book.active.append(['THC',20])
    buf = BytesIO(); book.save(buf)
    parser = Parser(); parser.parse(buf.getvalue(), 'synthetic.xlsx')
    doc = Document(); table = doc.add_table(rows=1, cols=2)
    table.cell(0,0).text='Doc Fee'; table.cell(0,1).text='15'
    buf = BytesIO(); doc.save(buf); parser.parse(buf.getvalue(), 'synthetic.docx')
    assert any('THC | 20' in e.text for e in parser.evidence)
    assert any('Doc Fee | 15' in e.text for e in parser.evidence)
