from io import BytesIO
from email.message import EmailMessage
import zipfile
import pytest
from billing_agent.parsers import Parser
from billing_agent.tabular import xlsx_rows
from billing_agent.models import Charge, Extraction
from billing_agent.normalization import normalize_extraction


def workbook(formula=False):
    import openpyxl
    book=openpyxl.Workbook()
    sheet=book.active
    sheet.merge_cells('A1:D1')
    sheet['A1']='Synthetic invoice'
    sheet.append(['Fee','Qty','Unit price','Amount'])
    sheet.append(['Handling Fee USD',2,20,'=B3*C3' if formula else 40])
    stream=BytesIO();book.save(stream)
    return stream.getvalue()


def with_cache(data):
    stream=BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as source,zipfile.ZipFile(stream,'w') as target:
        for name in source.namelist():
            raw=source.read(name)
            if name=='xl/worksheets/sheet1.xml':
                raw=raw.replace(b'<f>B3*C3</f><v></v>',b'<f>B3*C3</f><v>40</v>')
            target.writestr(name,raw)
    return stream.getvalue()


def test_excel_merge_keeps_anchor_once():
    rows=list(xlsx_rows(workbook()))
    assert rows[0][2].count('Synthetic invoice')==1
    assert rows[0][3]['cells'][1]['anchor']=='A1'
    assert rows[-1][2]=='Handling Fee USD | 2 | 20 | 40'


def test_named_table_columns_do_not_confuse_quantity_with_amount(monkeypatch):
    import openpyxl
    from billing_agent.extraction import extract_local
    book=openpyxl.Workbook();sheet=book.active
    sheet.append(['Description','Qty','Unit Price','Currency','Amount'])
    sheet.append(['Handling Fee',2,20,'USD',40])
    stream=BytesIO();book.save(stream)
    parser=Parser();parser.parse(stream.getvalue(),'bill.xlsx')
    row=parser.evidence[-1]
    assert row.structured['charge_cells']['amount']=='40'
    assert row.structured['business_cells']['quantity']['value']=='2'
    monkeypatch.setattr('urllib.request.OpenerDirector.open',lambda *args,**kwargs:(_ for _ in ()).throw(AssertionError('No model call for recognized row')))
    ext=normalize_extraction(extract_local([row]),parser.evidence)
    assert not ext.errors
    assert len(ext.charges)==1 and ext.charges[0].amount==40
    assert ext.charges[0].quantity==2 and ext.charges[0].unit_price==20


@pytest.mark.parametrize('cached',[False,True])
def test_formula_never_becomes_trusted_amount(cached):
    data=workbook(True)
    if cached:data=with_cache(data)
    parser=Parser();parser.parse(data,'bill.xlsx')
    row=parser.evidence[-1]
    assert '=B3*C3' not in row.text
    assert row.structured['cells'][-1]['formula']=='=B3*C3'
    assert row.structured['cells'][-1]['cached_value']==(40 if cached else None)
    c=Charge(charge_item='Handling Fee',amount=40,currency='USD',evidence_ids=[row.id],amount_quote='40')
    result=normalize_extraction(Extraction(charges=[c]),parser.evidence)
    assert any('formula' in w.lower() for w in result.charges[0].warnings)
    assert '[UNRESOLVED_FORMULA]' in row.text if not cached else '40' in row.text


def test_doc_converter_missing_is_actionable(monkeypatch):
    monkeypatch.setattr('billing_agent.conversion.soffice_path',lambda:(_ for _ in ()).throw(RuntimeError('DOC conversion requires LibreOffice')))
    parser=Parser();parser.parse(bytes.fromhex('d0cf11e0a1b11ae1')+b'fake','bill.doc')
    assert 'requires LibreOffice' in parser.issues[0]['error']


def test_doc_conversion_timeout_records_error_and_continues(monkeypatch):
    import subprocess
    monkeypatch.setattr('billing_agent.conversion.convert',lambda *args:(_ for _ in ()).throw(subprocess.TimeoutExpired('soffice',60)))
    parser=Parser();parser.parse(bytes.fromhex('d0cf11e0a1b11ae1')+b'fake','bill.doc')
    parser.parse(b'Handling Fee USD 40','other.txt')
    assert parser.issues and parser.evidence[-1].source=='other.txt'


def test_word_merge_and_body_order():
    from docx import Document
    doc=Document();doc.add_paragraph('Before table')
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).merge(table.cell(0,1)).text='Merged title'
    table.cell(1,0).text='Handling Fee USD';table.cell(1,1).text='40'
    doc.add_paragraph('After table')
    stream=BytesIO();doc.save(stream)
    parser=Parser();parser.parse(stream.getvalue(),'bill.docx')
    texts=[e.text for e in parser.evidence]
    assert texts[0]=='Before table' and texts[-1]=='After table'
    assert sum(t.count('Merged title') for t in texts)==1
    assert 'Handling Fee USD | 40' in texts


def test_html_mail_and_nested_word_attachment():
    from docx import Document
    doc=Document();doc.add_paragraph('Documentation Fee USD 10')
    stream=BytesIO();doc.save(stream)
    inner=EmailMessage();inner.set_content('Inside')
    inner.add_attachment(stream.getvalue(),maintype='application',subtype='octet-stream',filename='bill.docx')
    outer=EmailMessage();outer.set_content('<p>Handling Fee USD 40</p><script>USD 999</script>',subtype='html')
    outer.add_attachment(inner)
    parser=Parser();parser.parse(outer.as_bytes(),'outer.eml')
    assert any('::' in e.source and e.source.endswith('bill.docx') for e in parser.evidence)
    body=next(e for e in parser.evidence if e.source=='outer.eml' and e.location=='body')
    assert body.text=='Handling Fee USD 40'


def test_scanned_pdf_renders_and_preserves_page_identity(monkeypatch):
    from PIL import Image
    im=Image.new('RGB',(200,300),'white');stream=BytesIO()
    im.save(stream,format='PDF',save_all=True,append_images=[im])
    seen=[]
    def local(self,data,name,pdf_page=None,pdf_digest=None):
        seen.append((name,pdf_page,pdf_digest,data[:8]))
        self.add(name,f'page:{pdf_page}','Handling Fee USD 40','local_pdf_ocr_unverified')
    monkeypatch.setattr(Parser,'local_image',local)
    parser=Parser(pdf_mode='local');parser.parse(stream.getvalue(),'scan.pdf')
    assert [s[1] for s in seen]==[1,2]
    assert all(s[0]=='scan.pdf' and s[2] and s[3]==b'\x89PNG\r\n\x1a\n' for s in seen)
    assert all(e.method=='local_pdf_ocr_unverified' for e in parser.evidence)


def test_bad_spreadsheet_does_not_stop_other_files():
    parser=Parser();parser.parse(b'corrupt','bad.xlsx');parser.parse(b'Hello','okay.txt')
    assert 'error' in parser.issues[0] and parser.evidence[-1].source=='okay.txt'


def test_vertical_merge_is_not_silently_forward_filled():
    import openpyxl
    book=openpyxl.Workbook();s=book.active
    s.append(['Description','Currency','Amount'])
    s.append(['Handling Fee','USD',40]);s.append([None,'USD',20])
    s.merge_cells('A2:A3')
    stream=BytesIO();book.save(stream)
    rows=list(xlsx_rows(stream.getvalue()))
    assert rows[-1][3]['charge_cells'] is None
    assert rows[-1][3]['cells'][0]['anchor']=='A2'
    assert rows[-1][3]['review_reasons']


def test_excel_error_cell_is_explicitly_unsafe():
    import openpyxl
    book=openpyxl.Workbook();book.active.append(['Handling Fee USD','#DIV/0!'])
    stream=BytesIO();book.save(stream)
    parser=Parser();parser.parse(stream.getvalue(),'error.xlsx')
    assert '[CELL_ERROR]' in parser.evidence[0].text
    assert parser.evidence[0].structured['review_reasons']
