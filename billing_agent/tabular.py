"""Cell provenance, merged-cell anchors and explicit formula review policy."""
from io import BytesIO
import re
from decimal import Decimal, InvalidOperation


def business_cells(headers, values):
    aliases={'qty':'quantity','quantity':'quantity','数量':'quantity','unit price':'unit_price',
             'rate':'unit_price','单价':'unit_price','p/c':'payment_term','payment term':'payment_term'}
    result={}
    for header,value in zip(headers,values):
        field=aliases.get(str(header).strip().casefold())
        if not field or value is None:
            continue
        raw=str(value)
        if field=='payment_term':
            code={'P':'PP','C':'CC','PREPAID':'PP','COLLECT':'CC'}.get(raw.upper(),raw.upper())
            if code in {'PP','CC'}:result[field]={'value':code,'quote':raw}
        else:
            try:
                number=Decimal(raw.replace(',',''))
                if number.is_finite():result[field]={'value':str(number),'quote':raw}
            except InvalidOperation:
                pass
    return result


def is_header(values):
    labels={str(v).strip().casefold() for v in values}
    return bool(labels & {'description','fee','charge','费用名称'}) and bool(labels & {'amount','金额','currency','币种'})


def charge_cells(headers,values):
    """Only explicit named columns establish fee/amount boundaries."""
    if ambiguous_headers(headers):return None
    columns={str(h).strip().casefold():v for h,v in zip(headers,values)}
    name=next((columns[k] for k in ['description','fee','charge','费用名称'] if k in columns),None)
    amount=next((columns[k] for k in ['amount','金额'] if k in columns),None)
    currency=next((columns[k] for k in ['currency','币种'] if k in columns),None)
    try:
        number=Decimal(str(amount).replace(',',''))
        if not number.is_finite() or not name or not re.fullmatch(r'[A-Z]{3}',str(currency)):
            return None
    except InvalidOperation:
        return None
    return {'description':str(name),'amount':str(number),'currency':currency,'charge_code':None}


def ambiguous_headers(headers):
    labels=[str(h).strip().casefold() for h in headers]
    return any(sum(v in group for v in labels)>1 for group in
        [{'amount','金额'},{'currency','币种'},{'description','fee','charge','费用名称'}])


def xlsx_rows(data):
    import openpyxl
    book = openpyxl.load_workbook(BytesIO(data),data_only=False,read_only=False,keep_links=False)
    cached = openpyxl.load_workbook(BytesIO(data),data_only=True,read_only=False,keep_links=False)
    try:
        for sheet in book:
            headers=[]
            if sheet.max_row*sheet.max_column>200_000:
                raise ValueError('Spreadsheet cell budget exceeded')
            anchors = {}
            for merged in sheet.merged_cells.ranges:
                for row in sheet.iter_rows(min_row=merged.min_row,max_row=merged.max_row,
                                           min_col=merged.min_col,max_col=merged.max_col):
                    for cell in row:
                        anchors[cell.coordinate] = sheet.cell(merged.min_row,merged.min_col).coordinate
            for row in sheet.iter_rows():
                values=[]
                cells=[]
                reasons=[]
                for cell in row:
                    anchor=anchors.get(cell.coordinate,cell.coordinate)
                    origin=sheet[anchor]
                    # Never duplicate a vertically/horizontally merged fee cell.
                    value=origin.value if cell.coordinate==anchor else None
                    item={'cell':cell.coordinate,'anchor':anchor,'anchor_value':origin.value,'value':value}
                    if anchor!=cell.coordinate and origin.row!=cell.row and origin.value is not None:
                        reasons.append('Vertical merged cell continuation; row association requires review')
                    if cell.coordinate==anchor and origin.data_type=='f':
                        value=cached[sheet.title][anchor].value
                        item.update(formula=origin.value,cached_value=value,cache_verified=False)
                        reasons.append('Spreadsheet formula has no cached value' if value is None else 'Spreadsheet formula cache freshness unverified')
                        value='[UNRESOLVED_FORMULA]' if value is None else value
                    elif origin.data_type=='e':
                        reasons.append('Spreadsheet contains an error cell')
                        value='[CELL_ERROR]'
                    item['display']=value
                    cells.append(item)
                    values.append('' if value is None else str(value))
                if any(values):
                    if is_header(values):
                        headers=values
                    if ambiguous_headers(headers):reasons.append('Ambiguous repeated table columns; manual layout verification required')
                    display=[c['display'] for c in cells]
                    yield sheet.title,row[0].row,' | '.join(values),{'kind':'spreadsheet_row','cells':cells,
                        'charge_cells':charge_cells(headers,display),
                        'business_cells':business_cells(headers,[c['display'] for c in cells]),
                        'review_reasons':list(dict.fromkeys(reasons))}
    finally:
        book.close()
        cached.close()


def word_blocks(data):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc=Document(BytesIO(data))
    cell_count=0
    def table_rows(table,path,depth):
        nonlocal cell_count
        if depth>8:raise ValueError('Nested Word table depth limit exceeded')
        seen={};headers=[]
        for number,row in enumerate(table.rows,1):
            cell_count+=len(row.cells)
            if cell_count>200_000:raise ValueError('Word table cell budget exceeded')
            values=[];nested=[];cells=[]
            for column,cell in enumerate(row.cells,1):
                duplicate=cell._tc in seen
                values.append('' if duplicate else cell.text)
                cell_path=f'{path}/row:{number}/cell:{column}'
                cells.append({'column':column,'path':cell_path,'merged_duplicate':duplicate,
                    'origin_row':seen[cell._tc] if duplicate else number})
                if not duplicate:
                    seen[cell._tc]=number
                    nested.extend((child,f'{cell_path}/table:{i}') for i,child in enumerate(cell.tables,1))
            if is_header(values):headers=values
            reasons=[]
            if any(c['origin_row']!=number for c in cells):reasons.append('Vertical merged Word cell continuation; manual row association required')
            if ambiguous_headers(headers):reasons.append('Ambiguous repeated table columns; manual layout verification required')
            yield f'{path}/row:{number}',' | '.join(values) if any(v.strip() for v in values) else '',{'kind':'word_table_row',
                'table_path':path,'depth':depth,'row':number,'cells':cells,
                'review_reasons':reasons,
                'contains_nested_tables':bool(nested),'charge_cells':charge_cells(headers,values),
                'business_cells':business_cells(headers,values)}
            for child,child_path in nested:
                yield from table_rows(child,child_path,depth+1)
    for index,element in enumerate(doc.element.body):
        if element.tag.endswith('}p'):
            yield f'block:{index+1}/paragraph',Paragraph(element,doc).text,{}
        elif element.tag.endswith('}tbl'):
            yield from table_rows(Table(element,doc),f'block:{index+1}',0)


def html_text(html):
    from html.parser import HTMLParser
    class Text(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.parts=[]
            self.hidden=0
        def handle_starttag(self,tag,attrs):
            if tag in {'script','style'}:
                self.hidden+=1
            if tag in {'p','div','br','tr','li','blockquote'}:
                self.parts.append('\n')
            if tag in {'td','th'}:
                self.parts.append(' ')
        def handle_endtag(self,tag):
            if tag in {'script','style'}:
                self.hidden=max(0,self.hidden-1)
            if tag in {'p','div','tr','li','blockquote'}:
                self.parts.append('\n')
        def handle_data(self,text):
            if not self.hidden:
                self.parts.append(text)
    parser=Text()
    parser.feed(html)
    return '\n'.join(line.strip() for line in ''.join(parser.parts).splitlines() if line.strip())
